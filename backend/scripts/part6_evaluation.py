"""Part 6: full-suite evaluation, including the sweep-without-requery design
and the four-category (+2 verification-layer) routing scorecard.

Architecture note: this script talks to ValidationEngine / correlate_events /
DiagnosticAgent / evaluate_gate / attempt_fix directly, not through the
FastAPI /ingest endpoint. That's deliberate, not a shortcut: those ARE the
production functions ingest.py orchestrates (already covered end-to-end by
tests/test_action_executor_api.py), and calling them directly is what makes
"diagnose once, sweep thresholds without re-running the pipeline" possible -
going through HTTP/DB per threshold would re-invoke diagnosis (fine, it's
cached) but would also commit mutations (baselines, fix_chain) that would
contaminate later threshold iterations on "the same" corruption.

LLM PROVIDER NOTE: this script defaults to the real provider (get_llm_client(),
reading LLM_PROVIDER/GEMINI_API_KEY from env same as production - Gemini is
the only provider this deployment runs). This repository's own verification
run used a deterministic SmartFakeLLMClient instead, because no API key is
configured in this environment - that is disclosed in the printed output,
not hidden. Set GEMINI_API_KEY to run this for real.

Run: .venv/Scripts/python.exe scripts/part6_evaluation.py
"""

from __future__ import annotations

import json
import os
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

from dotenv import load_dotenv  # noqa: E402

load_dotenv()  # GEMINI_API_KEY etc. from .env - this script doesn't import app.main, so nothing else loads it

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

from app.connectors.file_connector import FileConnector  # noqa: E402
from app.contract import DataContract, SourceType  # noqa: E402
from app.correlation import CorrelatedGroup, correlate_events  # noqa: E402
from app.diagnosis.agent import DiagnosticAgent  # noqa: E402
from app.diagnosis.cache import DiagnosisCache  # noqa: E402
from app.diagnosis.models import CauseCategory, Diagnosis, FixAction, RiskLevel, SuggestedFix  # noqa: E402
from app.gate import AUTO_APPLY, build_fix_spec, evaluate_gate, permitted_actions  # noqa: E402
from app.llm.base import LLMClient  # noqa: E402
from app.profiling import BaselineProfiler  # noqa: E402
from app.repair import attempt_fix  # noqa: E402
from app.validation.engine import ValidationEngine  # noqa: E402
from tests.corruption import CorruptionSuite  # noqa: E402
from tests.corruption.injectors import change_dtype  # noqa: E402

CLEAN_DATA_SEED = 0
CORRUPTION_BASE_SEED = 100
COMPOSE_SEED = 200
FILE_CASE_SEED = 300
THRESHOLDS = [0.5, 0.55, 0.6, 0.65, 0.7, 0.75, 0.8, 0.85, 0.9, 0.95]

EXPECTED_CAUSE_CATEGORY = {
    "rename_column": "rename",
    "change_dtype": "dtype_change",
    "inject_nulls": "null_flood",
    "drop_column": "column_dropped",
    "shift_distribution": "distribution_shift",
    "inject_whitespace_case": "whitespace_case",
    "truncate_rows": "row_loss",
}

# Worse (lower) wins when a corruption's several groups disagree - the
# per-corruption figure is only as good as its weakest matched group, and
# a single wrongly-auto-fixed group makes the whole corruption count as one.
SEVERITY_RANK = {
    "wrongly_auto_fixed": 0,
    "wrongly_escalated": 1,
    "auto_fix_reverted_by_verification": 2,
    "correctly_escalated": 3,
    "correctly_auto_fixed": 4,
}


# =============================================================================
# A deterministic, ground-truth-aware fake LLM, used ONLY because this
# environment has no GEMINI_API_KEY. Confidence varies by corruption family
# (not a constant) so the threshold sweep has something real to show.
# =============================================================================


class SmartFakeLLMClient(LLMClient):
    temperature = 0.0
    model_name = "smart-fake-eval-model (no GEMINI_API_KEY in this environment)"

    def __init__(self):
        self.calls: list[str] = []

    @property
    def call_count(self) -> int:
        return len(self.calls)

    def complete(self, system: str, user: str, response_schema):
        self.calls.append(user)
        has_missing = "missing_column" in user
        has_unexpected = "unexpected_column" in user

        if has_missing and has_unexpected:
            return Diagnosis(
                cause_category=CauseCategory.RENAME,
                likely_cause="column relabeled upstream; data intact under the new name",
                suggested_fix=SuggestedFix(action=FixAction.RENAME_COLUMN),
                risk_level=RiskLevel.LOW,
                confidence=0.92,
            )
        if "dtype_mismatch" in user:
            return Diagnosis(
                cause_category=CauseCategory.DTYPE_CHANGE,
                likely_cause="numeric column arrived as text",
                suggested_fix=SuggestedFix(action=FixAction.SAFE_TYPE_CAST),
                risk_level=RiskLevel.LOW,
                confidence=0.85,
            )
        if "categorical_drift" in user:
            return Diagnosis(
                cause_category=CauseCategory.WHITESPACE_CASE,
                likely_cause="stray whitespace/case variance in category values",
                suggested_fix=SuggestedFix(action=FixAction.STRIP_WHITESPACE),
                risk_level=RiskLevel.LOW,
                confidence=0.78,
            )
        if has_missing:  # uncorrelated -> genuine drop
            return Diagnosis(
                cause_category=CauseCategory.COLUMN_DROPPED,
                likely_cause="column absent from current schema, no matching new column",
                suggested_fix=SuggestedFix(action=FixAction.ESCALATE),
                risk_level=RiskLevel.HIGH,
                confidence=0.88,
            )
        if "null_threshold" in user:
            return Diagnosis(
                cause_category=CauseCategory.NULL_FLOOD,
                likely_cause="null rate well above baseline tolerance",
                suggested_fix=SuggestedFix(action=FixAction.ESCALATE),
                risk_level=RiskLevel.HIGH,
                confidence=0.9,
            )
        if "distribution_drift" in user:
            return Diagnosis(
                cause_category=CauseCategory.DISTRIBUTION_SHIFT,
                likely_cause="numeric distribution has moved relative to baseline",
                suggested_fix=SuggestedFix(action=FixAction.ESCALATE),
                risk_level=RiskLevel.HIGH,
                confidence=0.87,
            )
        if "row_count_drop" in user:
            return Diagnosis(
                cause_category=CauseCategory.ROW_LOSS,
                likely_cause="row count dropped well below baseline",
                suggested_fix=SuggestedFix(action=FixAction.ESCALATE),
                risk_level=RiskLevel.HIGH,
                confidence=0.9,
            )
        return Diagnosis(
            cause_category=CauseCategory.UNKNOWN,
            likely_cause="no confident cause identified",
            suggested_fix=SuggestedFix(action=FixAction.ESCALATE),
            risk_level=RiskLevel.HIGH,
            confidence=0.5,
        )


# =============================================================================
# Fixtures
# =============================================================================


def make_clean_orders(n: int = 300, seed: int = CLEAN_DATA_SEED) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    cities = ["New York", "Los Angeles", "San Francisco", "Chicago", "Houston"]
    return pd.DataFrame(
        {
            "order_id": range(1, n + 1),
            "customer_city": rng.choice(cities, size=n),
            "price": rng.normal(loc=120.0, scale=35.0, size=n).round(2).clip(min=1.0),
            "quantity": rng.integers(1, 6, size=n),
        }
    )


def _write_excel_text_column(df: pd.DataFrame, column: str, path: Path) -> None:
    """Writes `column` as explicitly Text-formatted cells (data_type='s',
    number_format='@') - what a real 'numbers stored as text' Excel export
    looks like at the file-format level."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(list(df.columns))
    col_idx = list(df.columns).index(column) + 1
    for _, row in df.iterrows():
        ws.append(list(row))
        cell = ws.cell(row=ws.max_row, column=col_idx)
        cell.number_format = "@"
        cell.data_type = "s"
    wb.save(path)


def build_cases(clean_df: pd.DataFrame, scratch_dir: Path) -> list[tuple[str, DataContract, list]]:
    """Returns (case_name, contract, ground_truths) tuples: the single-fault
    suite, one composed multi-fault case, and the two file-format-specific
    dtype cases (CSV dirty-prefix / Excel clean-text)."""
    suite = CorruptionSuite()
    cases: list[tuple[str, DataContract, list]] = []

    for name, (df, truth) in suite.apply_all(clean_df, base_seed=CORRUPTION_BASE_SEED).items():
        cases.append((name, DataContract(data=df, source_type=SourceType.FILE, source_id="eval"), [truth]))

    composed_df, composed_truths = suite.compose(
        clean_df,
        [
            ("rename_column", {"column": "customer_city", "new_name": "city"}),
            ("shift_distribution", {"column": "price", "mode": "scale", "factor": 4.0}),
        ],
        seed=COMPOSE_SEED,
    )
    cases.append(
        ("composed_rename_and_shift", DataContract(data=composed_df, source_type=SourceType.FILE, source_id="eval"), composed_truths)
    )

    # CSV dirty-prefix: persists through a real CSV round trip, but safe_type_cast
    # correctly REJECTS the cast (a currency symbol is not a clean numeric string).
    csv_df, csv_truth = change_dtype(clean_df, seed=FILE_CASE_SEED, column="price", prefix="$")
    csv_path = scratch_dir / "change_dtype_csv_dirty.csv"
    csv_df.to_csv(csv_path, index=False)
    csv_contract = FileConnector(source_id="eval", file_path=str(csv_path)).fetch()
    cases.append(("change_dtype_csv_dirty_prefix", csv_contract, [csv_truth]))

    # Excel clean-text: a genuinely clean numeric string, stored as an explicitly
    # Text-formatted cell. This is safe_type_cast's actual happy path - CSV
    # structurally cannot provide one (see change_dtype's docstring). Reading it
    # back forces dtype=str for this column, bypassing FileConnector's default
    # pd.read_excel() call, which does NOT respect the cell's stored format
    # (see README "Phase 2 methodology notes" for the finding in full).
    xlsx_df, xlsx_truth = change_dtype(clean_df, seed=FILE_CASE_SEED + 1, column="price", prefix="")
    xlsx_path = scratch_dir / "change_dtype_excel_clean.xlsx"
    _write_excel_text_column(xlsx_df, "price", xlsx_path)
    xlsx_raw = pd.read_excel(xlsx_path, dtype={"price": str})
    xlsx_contract = DataContract(data=xlsx_raw, source_type=SourceType.FILE, source_id="eval")
    cases.append(("change_dtype_excel_clean_text", xlsx_contract, [xlsx_truth]))

    return cases


# =============================================================================
# Phase 1: diagnose once
# =============================================================================


@dataclass
class CollectedItem:
    case_name: str
    truth: object  # CorruptionGroundTruth
    group: CorrelatedGroup
    contract: DataContract
    diagnosis_json: dict
    diagnosis_source: str


def _match_truth(group: CorrelatedGroup, truths: list) -> object:
    if len(truths) == 1:
        return truths[0]
    group_columns = {m.column for m in group.members if m.column}
    for truth in truths:
        if group_columns & set(truth.target_columns):
            return truth
    return truths[0]


def diagnose_once(cases, baseline_profile, agent: DiagnosticAgent, engine: ValidationEngine) -> list[CollectedItem]:
    collected: list[CollectedItem] = []
    for case_name, contract, truths in cases:
        failures = engine.validate(contract, baseline_profile)
        groups = correlate_events(failures, baseline_profile, contract)
        for group in groups:
            outcome = agent.diagnose_group(group, baseline_profile, contract)
            collected.append(
                CollectedItem(
                    case_name=case_name,
                    truth=_match_truth(group, truths),
                    group=group,
                    contract=contract,
                    diagnosis_json=outcome.diagnosis_json,
                    diagnosis_source=outcome.source,
                )
            )
    return collected


# =============================================================================
# Phase 2: sweep (no new diagnosis calls - pure gate/repair replay)
# =============================================================================


def classify_at_threshold(item: CollectedItem, threshold: float, baseline_profile: dict, engine: ValidationEngine) -> dict:
    decision = evaluate_gate(item.group, item.diagnosis_json, threshold)
    diagnosed_cause = (item.diagnosis_json or {}).get("cause_category")
    expected_cause = EXPECTED_CAUSE_CATEGORY.get(item.truth.corruption_type)
    diagnosis_correct = diagnosed_cause == expected_cause

    if decision.decision == AUTO_APPLY:
        spec = build_fix_spec(item.group, decision.action)
        repair_result = attempt_fix(item.contract, item.group, decision.action, spec, baseline_profile, engine)
        if repair_result.verified:
            category = "correctly_auto_fixed" if item.truth.expected_risk_level == "low" else "wrongly_auto_fixed"
        else:
            category = "auto_fix_reverted_by_verification"
    else:
        category = "correctly_escalated" if item.truth.expected_risk_level == "high" else "wrongly_escalated"

    human_outcome = None
    if decision.decision != AUTO_APPLY:
        suggested = (item.diagnosis_json or {}).get("suggested_fix") or {}
        try:
            action = FixAction(suggested.get("action"))
        except ValueError:
            action = None
        if action is not None and action in permitted_actions(item.group):
            spec = build_fix_spec(item.group, action)
            human_repair = attempt_fix(item.contract, item.group, action, spec, baseline_profile, engine)
            human_outcome = "would_resolve" if human_repair.verified else "human_approved_but_verification_refused"
        else:
            human_outcome = "no_actionable_fix_suggested"

    return {
        "case": item.case_name,
        "corruption_type": item.truth.corruption_type,
        "n_events": len(item.group.members),
        "category": category,
        "diagnosis_correct": diagnosis_correct,
        "human_approval_outcome": human_outcome,
    }


def sweep(collected: list[CollectedItem], baseline_profile: dict, engine: ValidationEngine) -> dict[float, list[dict]]:
    return {t: [classify_at_threshold(item, t, baseline_profile, engine) for item in collected] for t in THRESHOLDS}


def confidence_distribution(collected: list[CollectedItem]) -> dict | None:
    """Phase 7.5 Part 5: min/max/mean/distinct-count over every diagnosis's
    self-reported confidence, from whichever provider actually ran (live or
    fake) - so calibration is something this script OBSERVES per run,
    never something assumed true because gate.py has a threshold to sweep
    against. None only when there is nothing to report at all (every
    diagnosis escalated with no confidence value)."""
    values = [item.diagnosis_json.get("confidence") for item in collected if item.diagnosis_json and item.diagnosis_json.get("confidence") is not None]
    if not values:
        return None
    distinct = sorted(set(values))
    return {
        "n": len(values),
        "min": min(values),
        "max": max(values),
        "mean": sum(values) / len(values),
        "distinct_values": distinct,
        "distinct_count": len(distinct),
    }


# =============================================================================
# Reporting
# =============================================================================


def _markdown_table(headers: list[str], rows: list[list]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    for row in rows:
        lines.append("| " + " | ".join(str(c) for c in row) + " |")
    return "\n".join(lines)


def main() -> None:
    # scratch_dir is genuinely ephemeral (the file-case CSV/xlsx fixtures
    # build_cases writes, deterministic from a fixed seed so regenerating
    # them is free) - the diagnosis CACHE is not, and previously lived
    # inside this same throwaway directory. That meant it worked WITHIN one
    # run and never across runs: every invocation got a fresh, empty
    # diagnoses.json and burned quota re-diagnosing the exact same
    # corruptions a prior run had already paid for (Phase 7.5 Part 4).
    # DiagnosisCache's own default (app/diagnosis/cache.py::DEFAULT_CACHE_PATH,
    # or DIAGNOSIS_CACHE_PATH from env) is a STABLE, gitignored path - the
    # same one production ingest uses - so passing no path override here at
    # all is what makes this cache actually persist across invocations.
    scratch_dir = Path(tempfile.mkdtemp(prefix="part6_eval_"))

    clean_df = make_clean_orders()
    baseline_profile = BaselineProfiler().profile(clean_df)
    cases = build_cases(clean_df, scratch_dir)

    using_fake = "GEMINI_API_KEY" not in os.environ and os.environ.get("LLM_PROVIDER", "gemini") == "gemini"
    if using_fake:
        llm_client = SmartFakeLLMClient()
        provider_note = "SmartFakeLLMClient (no GEMINI_API_KEY in this environment - NOT a live call)"
    else:
        from app.llm.factory import get_llm_client

        llm_client = get_llm_client()
        provider_note = f"live provider: {llm_client.model_name}"

    cache = DiagnosisCache()
    agent = DiagnosticAgent(llm_client=llm_client, cache=cache)
    engine = ValidationEngine()

    print("=== Config ===")
    print(f"provider: {provider_note}")
    print(f"model_name={llm_client.model_name} temperature={llm_client.temperature}")
    print(f"seeds: clean_data_seed={CLEAN_DATA_SEED} corruption_base_seed={CORRUPTION_BASE_SEED} "
          f"compose_seed={COMPOSE_SEED} file_case_seed={FILE_CASE_SEED}")
    print(f"thresholds swept: {THRESHOLDS}")
    print(f"diagnosis cache: {cache._path}")
    print()

    calls_before = getattr(llm_client, "call_count", None)
    collected = diagnose_once(cases, baseline_profile, agent, engine)
    calls_after = getattr(llm_client, "call_count", None)
    cache_hits = sum(1 for item in collected if item.diagnosis_source == "cache")
    llm_sourced = sum(1 for item in collected if item.diagnosis_source == "llm")
    escalated_failures = sum(1 for item in collected if item.diagnosis_source.startswith("escalated"))

    print("=== Quota usage (diagnose-once phase) ===")
    print(f"diagnosis groups: {len(collected)}")
    print(f"LLM calls made this run: {calls_after - calls_before if calls_before is not None else 'n/a'}")
    print(f"served from cache: {cache_hits}")
    print(f"served from live/fake LLM: {llm_sourced}")
    print(f"escalated due to LLM/parse failure: {escalated_failures}")
    print()

    print("=== Confidence distribution (Phase 7.5 Part 5 - calibration observed, not assumed) ===")
    confidence_stats = confidence_distribution(collected)
    if confidence_stats is None:
        print("no diagnoses with a confidence value to report")
    else:
        print(
            f"n={confidence_stats['n']} min={confidence_stats['min']:.3f} max={confidence_stats['max']:.3f} "
            f"mean={confidence_stats['mean']:.3f} distinct_values={confidence_stats['distinct_values']} "
            f"(count={confidence_stats['distinct_count']})"
        )
        if confidence_stats["distinct_count"] <= 3:
            print(
                "3 or fewer distinct confidence values across this many diagnoses - the threshold sweep below "
                "has little to no separating signal to work with; whatever it shows is a function of a handful of "
                "clustered values, not a spread the threshold is meaningfully discriminating within."
            )
    print()

    # Re-run diagnosis to prove the cache actually eliminates the second pass's calls.
    # call_count is a test-double-only concept (FakeLLMClient/SmartFakeLLMClient) -
    # a real provider client (GeminiClient) doesn't track it, so this falls back to
    # comparing diagnosis_source distributions instead of a raw count.
    calls_before_2 = getattr(llm_client, "call_count", None)
    second_pass = diagnose_once(cases, baseline_profile, agent, engine)
    calls_after_2 = getattr(llm_client, "call_count", None)
    if calls_before_2 is not None and calls_after_2 is not None:
        print(f"Re-running diagnose_once a second time made {calls_after_2 - calls_before_2} additional LLM calls "
              f"(should be 0 - everything was already cached).")
    else:
        second_pass_sources = {s: sum(1 for item in second_pass if item.diagnosis_source == s) for s in {"cache", "llm", "escalated_parse_failure", "escalated_quota_exhausted"}}
        print(f"call_count not tracked by {type(llm_client).__name__} (not a test double) - "
              f"second-pass diagnosis_source distribution instead: {second_pass_sources} "
              f"(a successful 'llm' diagnosis should become 'cache' on this second pass; an escalated one is never "
              f"cached, so it legitimately calls the provider again both times, which is not a cache bug).")
    print()

    print("=== Detection & diagnosis accuracy (threshold-independent) ===")
    detection_rows = []
    for item in collected:
        fired_family = item.group.members[0].rule_failed.split(":")[0]
        detected = item.truth.expected_detection == fired_family or any(
            m.rule_failed.startswith(item.truth.expected_detection) for m in item.group.members
        )
        detection_rows.append(
            [item.case_name, item.truth.corruption_type, item.truth.expected_detection, "yes" if detected else "no",
             (item.diagnosis_json or {}).get("cause_category"), EXPECTED_CAUSE_CATEGORY.get(item.truth.corruption_type),
             "yes" if (item.diagnosis_json or {}).get("cause_category") == EXPECTED_CAUSE_CATEGORY.get(item.truth.corruption_type) else "no"]
        )
    print(_markdown_table(
        ["case", "corruption_type", "expected_detection", "detected", "diagnosed_cause", "expected_cause", "diagnosis_correct"],
        detection_rows,
    ))
    print()

    sweep_results = sweep(collected, baseline_profile, engine)

    print("=== Routing accuracy by threshold (per-EVENT counts) ===")
    categories = list(SEVERITY_RANK) + []
    event_table_rows = []
    for cat in categories:
        row = [cat]
        for t in THRESHOLDS:
            n = sum(r["n_events"] for r in sweep_results[t] if r["category"] == cat)
            row.append(n)
        event_table_rows.append(row)
    print(_markdown_table(["category"] + [f"t={t}" for t in THRESHOLDS], event_table_rows))
    print()

    print("=== Routing accuracy by threshold (per-CORRUPTION counts; worst matched group wins) ===")
    corruption_table_rows = []
    for cat in categories:
        row = [cat]
        for t in THRESHOLDS:
            by_truth: dict[tuple, str] = {}
            for r in sweep_results[t]:
                key = (r["case"], r["corruption_type"])
                current = by_truth.get(key)
                if current is None or SEVERITY_RANK[r["category"]] < SEVERITY_RANK[current]:
                    by_truth[key] = r["category"]
            n = sum(1 for v in by_truth.values() if v == cat)
            row.append(n)
        corruption_table_rows.append(row)
    print(_markdown_table(["category"] + [f"t={t}" for t in THRESHOLDS], corruption_table_rows))
    print()

    print("*** WRONGLY AUTO-FIXED, prominently, per threshold (per-corruption) ***")
    for t in THRESHOLDS:
        by_truth: dict[tuple, str] = {}
        for r in sweep_results[t]:
            key = (r["case"], r["corruption_type"])
            current = by_truth.get(key)
            if current is None or SEVERITY_RANK[r["category"]] < SEVERITY_RANK[current]:
                by_truth[key] = r["category"]
        wrongly = [k for k, v in by_truth.items() if v == "wrongly_auto_fixed"]
        print(f"  t={t}: {len(wrongly)} wrongly-auto-fixed corruption(s) {wrongly if wrongly else ''}")
    print()

    print("=== Human-approval-if-escalated outcome, at default threshold t=0.8 (verification-layer check) ===")
    human_rows = [
        [r["case"], r["corruption_type"], r["category"], r["human_approval_outcome"]]
        for r in sweep_results[0.8]
        if r["human_approval_outcome"] is not None
    ]
    print(_markdown_table(["case", "corruption_type", "routing_category", "human_approval_outcome"], human_rows))
    print()

    print("Note: 'auto_fix_reverted_by_verification' and 'human_approved_but_verification_refused' are")
    print("successes of the verification layer, not failures - each is a case where the system attempted")
    print("or was asked to apply a fix that would NOT have resolved the issue, and caught it before committing.")


if __name__ == "__main__":
    main()
