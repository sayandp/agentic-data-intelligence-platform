"""The deck must render the WHOLE run, not a count of it.

The first version of the analytics section printed "abc_pareto: 5 finding(s)"
while 73 findings with full payloads sat persisted and unread - the numbers
existed, the deck just never looked at them. These tests hold the renderer to
showing what the run actually produced, and hold the export screen's preview
to promising the same shape the builder then renders.

Helpers come from test_export_pptx.py so both modules describe the same deck.
"""

from __future__ import annotations

from app.export.deck import build_deck, deck_outline

from tests.test_export_pptx import _deck_text, _slide_text, _sources


def _analytics_with_findings() -> dict:
    """One finding of every persisted type, with the payload keys the real
    analyses write (checked against a live run's /analytics response)."""
    return {
        "value_definition": {
            "monetary_column": "unit_price",
            "quantity_column": "qty",
            "derived": True,
            "label": "unit_price x qty",
            "note": "Value is `unit_price` x `qty` per row.",
        },
        "quantity_confirmation": {
            "role": "quantity",
            "reason": "`unit_price` reads as a unit price.",
            "candidates": [{"column": "qty", "reasons": ["contains negative values"]}],
        },
        "applicability": [
            {"analysis": "abc_pareto", "applicable": True, "missing_requirements": []},
            {"analysis": "market_basket_x", "applicable": False, "missing_requirements": ["needs a line-item identifier"]},
        ],
        "results": [
            {
                "analysis": "abc_pareto",
                "ran": True,
                "findings": [
                    {
                        "finding_type": "concentration_band",
                        "payload": {
                            "finding_type": "concentration_band",
                            "band": "A",
                            "entity_count": 901,
                            "entity_share": 0.16642,
                            "value_share": 0.799901,
                        },
                    },
                    {
                        "finding_type": "concentration_curve",
                        "payload": {
                            "finding_type": "concentration_curve",
                            "top_20_percent_entity_count": 1083,
                            "top_20_percent_value_share": 0.8331,
                            "concentration_floor": 0.5,
                            "concentration_is_weak": False,
                        },
                    },
                    {
                        "finding_type": "non_contributing_entities",
                        "payload": {
                            "finding_type": "non_contributing_entities",
                            "entity_count": 284,
                            "zero_net_count": 283,
                            "negative_net_count": 1,
                            "note": "Held out of the value ranking.",
                        },
                    },
                ],
            },
            {
                "analysis": "rfm",
                "ran": True,
                "findings": [
                    {
                        "finding_type": "segment_profile",
                        "payload": {
                            "finding_type": "segment_profile",
                            "segment": "champions",
                            "method": "rfm_rule",
                            "entity_count": 1357,
                            "entity_share": 0.228374,
                            "value_share": 0.598823,
                        },
                    }
                ],
            },
            {
                "analysis": "cohort_retention",
                "ran": True,
                "findings": [
                    {
                        "finding_type": "retention_matrix",
                        "payload": {
                            "finding_type": "retention_matrix",
                            "granularity": "month",
                            "cohort_labels": ["2009-12", "2010-01"],
                            "mean_retained_share": [1.0, 0.250179, 0.22951],
                        },
                    }
                ],
            },
            {
                "analysis": "market_basket",
                "ran": True,
                "findings": [
                    {
                        "finding_type": "association_rule",
                        "payload": {
                            "finding_type": "association_rule",
                            "antecedent": ["PINK TEACUP"],
                            "consequent": ["GREEN TEACUP"],
                            "support": 0.023092,
                            "confidence": 0.8153,
                            "lift": 22.0259,
                            "transaction_count": 49353,
                        },
                    }
                ],
            },
            {
                "analysis": "retention_churn",
                "ran": True,
                "findings": [
                    {
                        "finding_type": "repeat_behaviour",
                        "payload": {
                            "finding_type": "repeat_behaviour",
                            "entity_count": 5942,
                            "repeat_entity_count": 5796,
                            "repeat_rate": 0.975429,
                            "inactivity_window_days": 90,
                            "inactive_entity_count": 3020,
                            "inactive_share": 0.5082,
                            "gap_days_median": 12.0,
                        },
                    }
                ],
            },
            {
                "analysis": "historical_clv",
                "ran": True,
                "findings": [
                    {
                        "finding_type": "lifetime_value",
                        "payload": {
                            "finding_type": "lifetime_value",
                            "segment": "champions",
                            "entity_count": 1357,
                            "average_order_value": 3.3799,
                            "purchase_frequency": 395.7384,
                            "observed_lifespan_days": 544.091,
                            "historical_value_per_entity": 1337.5427,
                        },
                    }
                ],
            },
        ],
    }


# ---- analytics ----


def test_every_analysis_ships_its_RESULTS_not_a_count_of_them():
    text = _deck_text(build_deck(_sources(analytics=_analytics_with_findings())))

    assert "901 entities" in text and "80.0% of value" in text  # concentration band
    assert "1083" in text and "83.3%" in text  # concentration curve
    assert "champions" in text and "1357 entities" in text  # rfm segment
    assert "PINK TEACUP -> GREEN TEACUP" in text  # basket rule
    assert "22.0259" in text  # its lift
    assert "5796 of 5942 entities" in text  # repeat behaviour
    assert "1,337.5427 per entity" in text  # historical clv
    assert "finding(s)" not in text, "the deck is still reporting counts instead of results"


def test_each_applicable_analysis_gets_its_own_named_slide():
    text = _deck_text(build_deck(_sources(analytics=_analytics_with_findings())))

    for expected in [
        "Value concentration (ABC/Pareto)",
        "RFM segments",
        "Cohort retention",
        "Market basket rules",
        "Repeat behaviour and churn",
        "Historical customer value",
    ]:
        assert expected in text, f"no slide for {expected}"


def test_the_value_basis_and_its_open_question_reach_the_deck():
    """A value figure cannot be used without knowing what it measured."""
    text = _deck_text(build_deck(_sources(analytics=_analytics_with_findings())))

    assert "unit_price` x `qty" in text
    assert "A column role is still unconfirmed" in text
    assert "contains negative values" in text
    assert "283 netting to zero" in text


def test_a_confirmed_role_is_named_in_the_deck():
    confirmed = [{"role": "monetary", "column_name": "total_spend"}]
    text = _deck_text(build_deck(_sources(analytics=_analytics_with_findings(), confirmed_roles=confirmed)))

    assert "monetary = `total_spend`" in text


def test_a_run_with_no_confirmed_role_says_the_roles_were_detected():
    text = _deck_text(build_deck(_sources(analytics=_analytics_with_findings())))

    assert "No column role on this source was confirmed by a person" in text


def test_a_weak_pareto_concentration_is_reported_as_such():
    analytics = _analytics_with_findings()
    analytics["results"][0]["findings"][1]["payload"].update(
        concentration_is_weak=True, concentration_note="Value is spread more evenly than the 50% floor."
    )
    text = _deck_text(build_deck(_sources(analytics=analytics)))

    assert "spread more evenly" in text


def test_a_strong_concentration_still_reports_the_floor_it_was_measured_against():
    """Reported either way - a run that is NOT weakly concentrated is a
    finding, not an absence of one."""
    text = _deck_text(build_deck(_sources(analytics=_analytics_with_findings())))

    assert "50.0% concentration floor" in text


def test_retention_is_reported_as_a_curve_because_that_is_what_is_persisted():
    """`mean_retained_share` is a list BY PERIOD and index 0 is 1.0 by
    construction - collapsing it to one number reports nothing."""
    text = _deck_text(build_deck(_sources(analytics=_analytics_with_findings())))

    assert "+1 25.0%" in text
    assert "+2 23.0%" in text


def test_a_basket_rule_does_not_imply_its_support_count():
    """`transaction_count` is the basket set the rule was measured over, not
    the number of transactions carrying it."""
    text = _deck_text(build_deck(_sources(analytics=_analytics_with_findings())))

    assert "support 2.3% of 49,353 transactions" in text


def test_the_not_applicable_list_still_ships_alongside_the_results():
    text = _deck_text(build_deck(_sources(analytics=_analytics_with_findings())))

    assert "Analyses that did not apply" in text
    assert "market_basket_x" in text
    assert "needs a line-item identifier" in text


# ---- model ----


def test_the_model_slide_names_the_chosen_family():
    model = dict(_sources().model_runs[0])
    model.update(model_family="arima", row_count_trained_on=25, split_strategy="forward_chaining")
    text = _deck_text(build_deck(_sources(model_runs=[model])))

    assert "model chosen: arima" in text
    assert "forward_chaining" in text


def test_a_model_that_lost_to_its_baseline_says_so():
    """A score beside a baseline is only judgeable if the reader knows which
    direction is better, and a deck is read quickly."""
    model = dict(_sources().model_runs[0])
    model.update(out_of_sample_metric="rmse", out_of_sample_score=152220.17)
    text = _deck_text(build_deck(_sources(model_runs=[model])))

    assert "did not beat its baseline" in text


def test_a_model_that_beat_its_baseline_is_not_editorialised():
    model = dict(_sources().model_runs[0])
    model.update(out_of_sample_metric="rmse", out_of_sample_score=1.0)
    text = _deck_text(build_deck(_sources(model_runs=[model])))

    assert "did not beat its baseline" not in text


def test_an_unknown_metric_direction_produces_no_verdict():
    """Guessing the direction would invent a judgement the run never made."""
    model = dict(_sources().model_runs[0])
    model.update(out_of_sample_metric="custom_score", out_of_sample_score=0.5)
    text = _deck_text(build_deck(_sources(model_runs=[model])))

    assert "did not beat its baseline" not in text


def test_two_forecasts_of_the_same_column_are_distinguishable():
    first = dict(_sources().model_runs[0], question="forecast monthly Quantity")
    second = dict(_sources().model_runs[0], question="forecast weekly Quantity")
    text = _deck_text(build_deck(_sources(model_runs=[first, second])))

    assert "forecast monthly Quantity" in text
    assert "forecast weekly Quantity" in text


# ---- the preview cannot promise what the deck will not render ----


def test_the_outline_matches_the_deck_that_gets_built():
    sources = _sources(analytics=_analytics_with_findings())
    text = _deck_text(build_deck(sources))

    for entry in deck_outline(sources):
        # "Charts" names a GROUP of slides, each captioned with the chart's
        # own title, so it is the captions that prove it - not the word.
        if entry["section"] == "Charts":
            assert all(c["title"] in text for c in sources.chart_refs)
            continue
        assert entry["section"] in text, f"the outline promises {entry['section']!r}, which the deck does not contain"


def test_the_outline_calls_an_absent_section_a_placeholder_not_content():
    states = {e["section"]: e["state"] for e in deck_outline(_sources(model_runs=[], chart_refs=[], analytics=None))}

    assert states["Model results"] == "placeholder"
    assert states["Charts"] == "placeholder"
    # The one section that is never a placeholder and never dropped.
    assert states["Data quality context"] == "content"


def test_the_outline_lists_every_applicable_analysis_by_name():
    sections = [e["section"] for e in deck_outline(_sources(analytics=_analytics_with_findings()))]

    assert "Value concentration (ABC/Pareto)" in sections
    assert "Historical customer value" in sections


# ---- analytics charts reach the deck as images ----


def test_an_analysis_with_a_chart_gets_that_chart_as_a_slide_image():
    """The whole point of moving derivation to the backend: the deck had no
    spec to render, so these charts were simply absent from it."""
    from pptx import Presentation
    import io

    from app.analytics.chart_specs import charts_for_results

    analytics = _analytics_with_findings()
    analytics["charts"] = charts_for_results(analytics["results"])
    assert analytics["charts"], "the fixture produces no charts to test with"

    prs = Presentation(io.BytesIO(build_deck(_sources(analytics=analytics))))
    titles = {c["title"] for c in analytics["charts"]}

    for slide in prs.slides:
        text = _slide_text(slide)
        if text.strip() in titles:
            # 13 is PICTURE in python-pptx's shape-type enum.
            assert any(shape.shape_type == 13 for shape in slide.shapes), f"{text!r} slide carries no image"
            titles.discard(text.strip())

    assert not titles, f"no slide at all for {titles}"


def test_a_run_analysed_before_chart_specs_existed_still_gets_its_charts():
    """Backfill: old rows carry findings but no `charts`. They are derived on
    read by the same function, so the deck is not silently chart-less for
    exactly the runs that predate the feature."""
    from app.export.pipeline import _analytics_payload

    class _Row:
        findings_json = {"results": _analytics_with_findings()["results"]}

    payload = _analytics_payload(_Row())

    # market_basket, retention_churn and historical_clv deliberately have no
    # chart - they read better as tables.
    assert {c["analysis"] for c in payload["charts"]} == {"abc_pareto", "rfm", "cohort_retention"}
