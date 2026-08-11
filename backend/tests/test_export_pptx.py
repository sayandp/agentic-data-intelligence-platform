"""The deck is a RENDERER over persisted artifacts.

These tests hold it to that: nothing appears that the run did not already
produce, the quality context is always second, a section with no content
states why rather than vanishing, and the chrome obeys the same causal-
vocabulary ban the generated prose does.
"""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone

import pytest
from pptx import Presentation

from app.export.deck import DeckSources, build_deck
from app.narrative.config import DEFAULT_CAUSAL_LEXICON


def _sources(**overrides) -> DeckSources:
    base = dict(
        run_number=42,
        run_id="run-uuid",
        run_status="completed",
        source_label="orders.csv",
        started_at=datetime(2026, 8, 11, 9, 30, tzinfo=timezone.utc),
        generation_mode="llm",
        narrative_text=(
            "## Data Quality Context\n"
            "- The active baseline for this source is still PROVISIONAL.\n\n"
            "Revenue correlates with cost across 5,497 observations.\n\n"
            "Band A holds 3,253 entities.\n"
        ),
        grounded_claims=[{"claim_id": "claim-1", "claim_text": "Band A holds 3,253 entities.", "finding_ids": ["f-1"]}],
        chart_refs=[
            {
                "chart_id": "c1",
                "title": "revenue - distribution",
                "figure_json": {"data": [{"type": "bar", "x": [1, 2, 3], "y": [3, 2, 1]}], "layout": {}},
            }
        ],
        analytics={
            "value_definition": {"note": "Value taken directly from `revenue`."},
            "results": [
                {"analysis": "abc_pareto", "ran": True, "findings": [{"id": "a-1"}]},
                {"analysis": "market_basket", "ran": False, "findings": []},
            ],
            "applicability": [
                {"analysis": "abc_pareto", "applicable": True, "missing_requirements": []},
                {
                    "analysis": "market_basket",
                    "applicable": False,
                    "missing_requirements": ["needs a line-item identifier that groups within a transaction; none detected"],
                },
            ],
        },
        model_runs=[
            {
                "question": "forecast monthly revenue",
                "target_column": "revenue",
                "task_type": "forecast",
                "out_of_sample_metric": "rmse",
                "out_of_sample_score": 91234.5,
                "baseline_scores": [
                    {"model_family": "naive", "metric": "rmse", "out_of_sample_score": 142621.9},
                    {"model_family": "seasonal_naive", "metric": "rmse", "out_of_sample_score": 100048.2},
                ],
                "prediction_interval": {"lower": 121275.7, "upper": 498375.2, "confidence_level": 0.8},
                "excluded_features": [{"feature": "order_id", "reason": "identifier-like, one value per row"}],
                "state": "answered",
                "escalation_reason": None,
            }
        ],
        trace=[{"node": "ingestion", "edge_taken": "validate"}, {"node": "narrative", "edge_taken": None}],
        validation_events=[
            {"rule_failed": "null_threshold", "state": "resolved", "action_taken": "fill_nulls", "reversal": None, "resolved_by": "analyst"}
        ],
    )
    base.update(overrides)
    return DeckSources(**base)


def _slides(data: bytes):
    return list(Presentation(io.BytesIO(data)).slides)


def _slide_text(slide) -> str:
    return "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)


def _deck_text(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    parts = []
    for slide in prs.slides:
        parts.append(_slide_text(slide))
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


# ---- structure ----


def test_a_full_run_produces_every_section():
    slides = _slides(build_deck(_sources()))
    text = "\n".join(_slide_text(s) for s in slides)

    assert len(slides) >= 7
    for expected in [
        "orders.csv",                      # title
        "Data quality context",
        "Narrative",
        "revenue - distribution",          # a chart slide, captioned
        "Business analytics",
        "Analyses that did not apply",
        "Model results",
        "Audit trail",
    ]:
        assert expected in text, f"deck is missing the {expected!r} section"


@pytest.mark.parametrize(
    "overrides",
    [
        {},
        {"model_runs": []},
        {"chart_refs": []},
        {"analytics": None},
        {"narrative_text": None, "grounded_claims": []},
        {"generation_mode": "template"},
        {"run_status": "awaiting_approval"},
    ],
    ids=["full", "no-model", "no-charts", "no-analytics", "no-narrative", "template", "paused"],
)
def test_quality_context_is_always_the_second_slide(overrides):
    """The one slide that must never be an appendix and never be dropped.

    A fluent report read as authoritative over repaired or unvalidated data
    is the most harmful thing this system can hand someone, and a deck
    travels further from its context than a screen does.
    """
    slides = _slides(build_deck(_sources(**overrides)))

    assert len(slides) >= 2
    assert "Data quality context" in _slide_text(slides[1])


def test_a_run_with_no_model_gets_an_explanatory_slide_not_a_dropped_section():
    """Dropping the section would let a reader assume modelling was never
    attempted, when it may have been attempted and honestly refused."""
    text = "\n".join(_slide_text(s) for s in _slides(build_deck(_sources(model_runs=[]))))

    assert "Model results" in text
    assert "No forecast or model was requested" in text


def test_an_escalated_model_reports_the_escalation_rather_than_a_score():
    escalated = dict(_sources().model_runs[0])
    escalated.update(state="escalated", escalation_reason="no candidate beat the baseline")
    text = "\n".join(_slide_text(s) for s in _slides(build_deck(_sources(model_runs=[escalated]))))

    assert "Escalated" in text
    assert "no candidate beat the baseline" in text


def test_a_model_score_always_appears_next_to_its_baseline():
    """A score with no baseline beside it cannot be judged."""
    text = "\n".join(_slide_text(s) for s in _slides(build_deck(_sources())))

    assert "out-of-sample rmse" in text
    assert "naive=" in text and "seasonal_naive=" in text
    assert "prediction interval" in text
    assert "order_id" in text and "identifier-like" in text


# ---- the honesty of the analytics feature survives the export ----


def test_the_not_applicable_list_ships_in_the_deck_with_its_reasons():
    text = "\n".join(_slide_text(s) for s in _slides(build_deck(_sources())))

    assert "Analyses that did not apply" in text
    assert "market_basket" in text
    assert "needs a line-item identifier" in text


def test_a_template_generated_report_exports_as_a_labelled_template_deck():
    text = "\n".join(_slide_text(s) for s in _slides(build_deck(_sources(generation_mode="template"))))

    assert "narrative: template" in text
    assert "deterministic template" in text


# ---- the constraints the UI is held to ----


def test_no_causal_vocabulary_anywhere_in_a_generated_deck():
    """The same lexicon app/narrative/postchecks.py enforces on generated
    prose. Deck chrome is held to the standard the content is."""
    text = _deck_text(build_deck(_sources()))

    for pattern in DEFAULT_CAUSAL_LEXICON:
        assert not re.search(pattern, text, re.I), f"causal vocabulary {pattern!r} reached the deck"


def test_grounded_claims_are_speaker_notes_not_body_text():
    """Each claim is the traceable evidence behind a sentence on the slide -
    a presenter needs them to hand, an audience does not need them read
    aloud."""
    prs = Presentation(io.BytesIO(build_deck(_sources())))
    notes = "\n".join(s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide)

    assert "claim-1" in notes
    assert "cites f-1" in notes


def test_the_deck_states_when_a_run_did_not_finish():
    text = "\n".join(_slide_text(s) for s in _slides(build_deck(_sources(run_status="awaiting_approval"))))

    assert "awaiting_approval" in text and "did not finish" in text


def test_every_empty_section_says_why_rather_than_disappearing():
    empty = _sources(chart_refs=[], analytics=None, model_runs=[], narrative_text=None, grounded_claims=[], validation_events=[], trace=[])
    text = "\n".join(_slide_text(s) for s in _slides(build_deck(empty)))

    assert "no charts" in text
    assert "No business analyses were computed" in text
    assert "No forecast or model was requested" in text
    assert "no narrative report" in text
    assert "No agent trace" in text


# ---- the deck cannot disagree with the dashboard ----


def test_charts_come_from_the_REPORTS_OWN_figure_json_not_a_second_spec():
    """One spec, produced once by app/narrative/charts.py. A slide and a
    screen showing different things for the same run would make both
    untrustworthy, so the deck renders the exact object the browser gets -
    including a title the backend already annotated."""
    from app.export.chart_images import _apply_deck_colours

    figure = {
        "data": [{"type": "scatter", "mode": "markers", "x": [1, 2], "y": [3, 4]}],
        "layout": {"title": "Price vs Quantity (every 213th of 1,067,371 points)"},
    }
    charts = [{"chart_id": "c1", "title": figure["layout"]["title"], "figure_json": figure}]
    slides = _slides(build_deck(_sources(chart_refs=charts)))
    text = "\n".join(_slide_text(s) for s in slides)

    # The caption is the backend's title verbatim, sampling note and all.
    assert "Price vs Quantity (every 213th of 1,067,371 points)" in text
    # And the data is passed through untouched - only colour is added.
    coloured = _apply_deck_colours(figure)
    assert coloured["data"][0]["x"] == [1, 2]
    assert coloured["data"][0]["y"] == [3, 4]
    assert coloured["data"][0]["type"] == "scatter"


def test_a_single_series_deck_chart_takes_one_colour_like_the_dashboard():
    """The same rule frontend/src/lib/plotly.ts applies. A lone bar chart
    carries its comparison in the bar lengths; colouring each bar
    differently would add nothing and imply the categories differ in kind."""
    from app.export.chart_images import DECK_ACCENT, DECK_CATEGORICAL, _apply_deck_colours

    one = _apply_deck_colours({"data": [{"type": "bar"}], "layout": {}})
    many = _apply_deck_colours({"data": [{"type": "bar"}, {"type": "bar"}, {"type": "bar"}], "layout": {}})

    assert one["data"][0]["marker"]["color"] == DECK_ACCENT
    assert [t["marker"]["color"] for t in many["data"]] == DECK_CATEGORICAL[:3]


def test_no_deck_colour_is_a_status_colour():
    """Same structural guarantee the UI has: the chart layer cannot reach
    the status vocabulary, so colour can never assert a verdict."""
    from app.export.chart_images import DECK_ACCENT, DECK_CATEGORICAL

    status = {"#157D3C", "#15803D", "#B91C1C", "#956400", "#9A6700", "#1E40AF"}
    for colour in [DECK_ACCENT, *DECK_CATEGORICAL]:
        assert colour.upper() not in status, f"{colour} is a status colour"
